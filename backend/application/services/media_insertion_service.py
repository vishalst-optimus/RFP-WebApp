"""
Media insertion service for handling image and document insertion into RFP sections
"""
import uuid
import html
import requests
import mimetypes
import io
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple
from docx import Document
from PyPDF2 import PdfReader
from pptx import Presentation
from application.services.cosmos_service.rfp_data_service import RFPDataService
from application.services.blob_service.azure_blob_service import (
    get_blob_service_client,
    container_exists_async,
    create_container_async
)
from core.models.media_schemas import (
    InsertMediaRequest, 
    InsertMediaResponse, 
    InsertMediaResponseData,
    MediaItemResponse,
    LocationData,
    DisplayInfo,
    ErrorInfo,
    ErrorDetails
)

class MediaInsertionService:
    def __init__(self):
        self.rfp_data_service = RFPDataService()
        self.blob_container_name = "sharepoint-files"
    
    def _convert_google_drive_url_to_download_url(self, sharing_url: str) -> str:
        """
        Convert a Google Drive sharing URL to a direct download URL
        Supports both file and document formats:
        - https://drive.google.com/file/d/FILE_ID/view?usp=drivesdk
        - https://docs.google.com/document/d/FILE_ID/edit?usp=drivesdk
        """
        try:
            file_id = None
            
            # Handle Google Drive file URLs
            if '/file/d/' in sharing_url:
                file_id_start = sharing_url.find('/file/d/') + 8
                file_id_end = sharing_url.find('/', file_id_start)
                if file_id_end == -1:
                    file_id_end = sharing_url.find('?', file_id_start)
                if file_id_end == -1:
                    file_id_end = len(sharing_url)
                file_id = sharing_url[file_id_start:file_id_end]
                
            # Handle Google Docs URLs
            elif '/document/d/' in sharing_url:
                file_id_start = sharing_url.find('/document/d/') + 12
                file_id_end = sharing_url.find('/', file_id_start)
                if file_id_end == -1:
                    file_id_end = sharing_url.find('?', file_id_start)
                if file_id_end == -1:
                    file_id_end = len(sharing_url)
                file_id = sharing_url[file_id_start:file_id_end]
                
            # Handle Google Sheets URLs
            elif '/spreadsheets/d/' in sharing_url:
                file_id_start = sharing_url.find('/spreadsheets/d/') + 16
                file_id_end = sharing_url.find('/', file_id_start)
                if file_id_end == -1:
                    file_id_end = sharing_url.find('?', file_id_start)
                if file_id_end == -1:
                    file_id_end = len(sharing_url)
                file_id = sharing_url[file_id_start:file_id_end]
                
            # Handle Google Slides URLs
            elif '/presentation/d/' in sharing_url:
                file_id_start = sharing_url.find('/presentation/d/') + 16
                file_id_end = sharing_url.find('/', file_id_start)
                if file_id_end == -1:
                    file_id_end = sharing_url.find('?', file_id_start)
                if file_id_end == -1:
                    file_id_end = len(sharing_url)
                file_id = sharing_url[file_id_start:file_id_end]
            
            if file_id:
                # Create direct download URL
                download_url = f"https://drive.google.com/uc?export=download&id={file_id}"
                print(f"🔄 Converted Google Drive/Docs URL: {sharing_url} -> {download_url}")
                return download_url
            else:
                print(f"⚠️ Unrecognized Google Drive URL format: {sharing_url}")
                return sharing_url
                
        except Exception as e:
            print(f"❌ Error converting Google Drive URL: {str(e)}")
            return sharing_url

    def _download_google_drive_file(self, file_id: str) -> requests.Response:
        """
        Download a file from Google Drive with proper session handling for large files
        """
        session = requests.Session()
        
        # First try the direct download URL
        URL = f"https://drive.google.com/uc?export=download&id={file_id}"
        response = session.get(URL, stream=True)
        
        # Check if we got HTML content (download warning page)
        content_type = response.headers.get('content-type', '').lower()
        if 'text/html' in content_type and response.status_code == 200:
            print(f"🔍 Received HTML response, checking for download warning page...")
            
            # Check if we got a download warning (for large files or virus scan)
            response_text = response.text
            if any(keyword in response_text.lower() for keyword in ['download_warning', 'virus scan warning', 'download anyway']):
                print(f"📋 Found download warning page, looking for confirm token...")
                
                # Look for the confirm token in the response
                import re
                # Try different patterns for the confirm token
                patterns = [
                    r'name="confirm"\s+value="([^"]+)"',
                    r'"downloadUrl":"[^"]*confirm=([^&"]+)',
                    r'confirm=([^&"]+)',
                    r'"confirm"\s*:\s*"([^"]+)"'
                ]
                
                confirm_token = None
                for pattern in patterns:
                    token_match = re.search(pattern, response_text)
                    if token_match:
                        confirm_token = token_match.group(1)
                        print(f"✅ Found confirm token: {confirm_token[:10]}...")
                        break
                
                if confirm_token:
                    # Make the actual download request with the confirm token
                    URL = f"https://drive.google.com/uc?export=download&confirm={confirm_token}&id={file_id}"
                    response = session.get(URL, stream=True)
                    print(f"🔄 Retried download with confirm token, status: {response.status_code}")
                else:
                    print(f"❌ Could not find confirm token in response")
            else:
                print(f"⚠️ Received HTML but not a download warning page. File might not be publicly accessible.")
                print(f"📄 Response preview: {response_text[:300]}...")
        
        return response
    
    async def _download_and_save_to_blob(
        self, 
        media_url: str, 
        media_name: str, 
        tenant_id: str, 
        user_id: str,
        rfp_name: Optional[str] = None,
        extract_content: bool = False,
        access_token: Optional[str] = None,
        source: str = 'sharepoint',
        media_type: str = 'document'
    ) -> Tuple[str, Optional[str]]:
        """
        Download media from various sources (SharePoint, Google Drive, etc.) and save to Azure Blob Storage
        Returns the blob URL and optionally extracted content for documents
        """
        try:
            # Prepare headers for the request
            headers = {}
            
            # Handle different sources
            if source == 'drive':
                # Extract file ID and use custom Google Drive download method
                file_id = None
                
                # Try to extract file ID from various Google URL formats
                if '/file/d/' in media_url:
                    file_id_start = media_url.find('/file/d/') + 8
                elif '/document/d/' in media_url:
                    file_id_start = media_url.find('/document/d/') + 12
                elif '/spreadsheets/d/' in media_url:
                    file_id_start = media_url.find('/spreadsheets/d/') + 16
                elif '/presentation/d/' in media_url:
                    file_id_start = media_url.find('/presentation/d/') + 16
                
                if 'file_id_start' in locals():
                    file_id_end = media_url.find('/', file_id_start)
                    if file_id_end == -1:
                        file_id_end = media_url.find('?', file_id_start)
                    if file_id_end == -1:
                        file_id_end = len(media_url)
                    
                    file_id = media_url[file_id_start:file_id_end]
                
                if file_id:
                    print(f"🌐 Downloading from Google Drive with file ID: {file_id}")
                    # Use custom Google Drive download method
                    response = self._download_google_drive_file(file_id)
                else:
                    # Fallback to direct URL if format not recognized
                    download_url = self._convert_google_drive_url_to_download_url(media_url)
                    response = requests.get(download_url, headers=headers, stream=True, timeout=30)
                    
            elif 'graph.microsoft.com' in media_url and access_token:
                download_url = media_url
                headers['Authorization'] = f'Bearer {access_token}'
                print(f"🔐 Using Bearer token for Graph API request")
                
                # Debug: Try to decode the token to check its properties (for debugging only)
                try:
                    import jwt
                    decoded = jwt.decode(access_token, options={"verify_signature": False})
                    print(f"🔍 Token audience (aud): {decoded.get('aud', 'Not found')}")
                    print(f"🔍 Token scopes (scp): {decoded.get('scp', 'Not found')}")
                    print(f"🔍 Token app ID (appid): {decoded.get('appid', 'Not found')}")
                except Exception as debug_e:
                    print(f"🔍 Could not decode token for debugging: {debug_e}")
                    
                print(f"🌐 Making request to: {download_url}")
                response = requests.get(download_url, headers=headers, stream=True, timeout=30)
                
            elif 'graph.microsoft.com' in media_url:
                download_url = media_url
                print(f"⚠️ Graph API URL detected but no access token provided")
                print(f"🌐 Making request to: {download_url}")
                response = requests.get(download_url, headers=headers, stream=True, timeout=30)
            else:
                # Default case for other sources
                download_url = media_url
                print(f"🌐 Making request to: {download_url}")
                response = requests.get(download_url, headers=headers, stream=True, timeout=30)
            
            # Debug response before raising for status
            print(f"📊 Response status: {response.status_code}")
            print(f"📊 Content-Type: {response.headers.get('content-type', 'Not specified')}")
            print(f"📊 Content-Length: {response.headers.get('content-length', 'Not specified')}")
            
            if response.status_code != 200:
                print(f"📊 Response headers: {dict(response.headers)}")
                print(f"📊 Response body: {response.text[:500]}...")
            
            # Special handling for Google Drive responses
            if source == 'drive':
                content_type = response.headers.get('content-type', '').lower()
                if 'text/html' in content_type:
                    print(f"⚠️ Warning: Google Drive returned HTML content instead of file")
                    print(f"💡 This usually means the file is not publicly accessible")
                    print(f"🔧 To fix: Right-click file in Google Drive > Share > Change to 'Anyone with the link'")
                    # Continue anyway - the HTML might still be useful for some purposes
                
            response.raise_for_status()
            
            # Get file extension from media name or try to determine from content-type
            file_extension = Path(media_name).suffix
            if not file_extension:
                content_type = response.headers.get('content-type', '')
                if content_type:
                    extension = mimetypes.guess_extension(content_type.split(';')[0])
                    if extension:
                        file_extension = extension
                        media_name = f"{Path(media_name).stem}{file_extension}"
            
            # Use the original filename (no GUID)
            clean_filename = media_name

            # Determine source subfolder based on source parameter
            source_folder = "sharepoint" if source == "sharepoint" else "drive"
            
            # Determine folder name based on media type (pluralized)
            media_folder = f"{media_type}s" if not media_type.endswith('s') else media_type

            # Create blob path: sharepoint-files/tenant_id/user_id/rfp_name/source_folder/media_type/filename
            if rfp_name:
                blob_name = f"{tenant_id}/{user_id}/{rfp_name}/{source_folder}/{media_folder}/{clean_filename}"
            else:
                blob_name = f"{tenant_id}/{user_id}/{source_folder}/{media_folder}/{clean_filename}"
            
            # Ensure container exists
            if not await container_exists_async(self.blob_container_name):
                await create_container_async(self.blob_container_name)
            
            # Get blob service client and upload
            blob_service_client = await get_blob_service_client()
            container_client = blob_service_client.get_container_client(self.blob_container_name)
            blob_client = container_client.get_blob_client(blob_name)
            
            # Upload the file content
            await blob_client.upload_blob(response.content, overwrite=True)
            
            # Extract document content if requested
            extracted_content = None
            if extract_content:
                content_type = response.headers.get('content-type', '')
                print(f"🔍 Attempting to extract content from document: {media_name}, content-type: {content_type}")
                extracted_content = self._extract_document_content(
                    response.content, media_name, content_type
                )
                if extracted_content:
                    print(f"📄 Extracted {len(extracted_content)} characters from document")
                else:
                    print(f"❌ Failed to extract content from document")
                print(f"📄 Content preview: {extracted_content[:200]}..." if extracted_content else "No content extracted")
            
            # Return the blob URL and extracted content
            blob_url = blob_client.url
            print(f"✅ Successfully saved media '{media_name}' to blob storage: {blob_url}")
            return blob_url, extracted_content
            
        except Exception as e:
            print(f"❌ Error downloading and saving media '{media_name}': {str(e)}")
            
            # Provide specific error message for authentication issues
            if '401' in str(e) or 'Unauthorized' in str(e):
                print(f"🔒 Authentication error - may need valid access token for Graph API")
                if 'graph.microsoft.com' in media_url:
                    print(f"💡 This appears to be a Graph API URL that requires authentication")
            elif source == 'drive' and ('403' in str(e) or 'Forbidden' in str(e)):
                print(f"🔒 Access denied for Google Drive file - file may not be publicly accessible")
                print(f"💡 Make sure the Google Drive file is shared with 'Anyone with the link'")
            
            # Return original URL as fallback
            return media_url, None
    
    def _extract_document_content(self, file_content: bytes, filename: str, mime_type: str) -> str:
        """
        Extract text content from various document formats
        """
        print(f"🔍 Extracting content from {filename}, mime-type: {mime_type}, size: {len(file_content)} bytes")
        
        try:
            file_extension = Path(filename).suffix.lower()
            print(f"🔍 File extension: {file_extension}")
            
            # First, check if we actually received file content or HTML
            if self._is_html_content(file_content):
                print(f"⚠️ Received HTML content instead of actual file - likely from Google Drive preview")
                print(f"💡 File sharing permissions may need to be updated to 'Anyone with the link'")
                
                # Try to extract any useful text from the HTML
                html_text = self._extract_text_from_html(file_content)
                if html_text and len(html_text.strip()) > 20:
                    return f"[HTML Preview Content from {filename}]\n\n{html_text}"
                else:
                    return f"[Document: {filename}] - Could not access file content. File may not be publicly accessible."
            
            if file_extension in ['.docx', '.doc'] or 'word' in mime_type.lower():
                print(f"📝 Processing as Word document")
                return self._extract_docx_content(file_content)
            elif file_extension == '.pdf' or 'pdf' in mime_type.lower():
                print(f"📄 Processing as PDF document")
                return self._extract_pdf_content(file_content)
            elif file_extension in ['.pptx', '.ppt'] or 'presentation' in mime_type.lower():
                print(f"📊 Processing as PowerPoint document")
                return self._extract_pptx_content(file_content)
            elif file_extension == '.txt' or 'text' in mime_type.lower():
                print(f"📝 Processing as text file")
                return file_content.decode('utf-8', errors='ignore')
            else:
                print(f"❌ Unsupported document type: {file_extension} / {mime_type}")
                return f"[Document: {filename}] - Content extraction not supported for this file type"
                
        except Exception as e:
            print(f"❌ Error extracting content from {filename}: {str(e)}")
            import traceback
            traceback.print_exc()
            return f"[Document: {filename}] - Error extracting content: {str(e)}"
    
    def _is_html_content(self, file_content: bytes) -> bool:
        """
        Check if the content is HTML (common when Google Drive returns preview instead of file)
        """
        try:
            # Check first 1000 bytes for HTML indicators
            content_preview = file_content[:1000].decode('utf-8', errors='ignore').lower()
            html_indicators = ['<html', '<head', '<body', '<!doctype html', '<div', '<span']
            return any(indicator in content_preview for indicator in html_indicators)
        except:
            return False
    
    def _extract_text_from_html(self, file_content: bytes) -> str:
        """
        Extract readable text from HTML content using BeautifulSoup
        """
        try:
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(file_content, 'html.parser')
            
            # Remove script and style elements
            for script in soup(["script", "style"]):
                script.decompose()
            
            # Get text and clean it up
            text = soup.get_text()
            lines = (line.strip() for line in text.splitlines())
            chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
            text = '\n'.join(chunk for chunk in chunks if chunk)
            
            return text[:2000]  # Limit to first 2000 characters
        except Exception as e:
            print(f"❌ Error extracting text from HTML: {str(e)}")
            return ""
    
    def _extract_text_from_html(self, file_content: bytes) -> str:
        """
        Extract readable text from HTML content using BeautifulSoup
        """
        try:
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(file_content, 'html.parser')
            
            # Remove script and style elements
            for script in soup(["script", "style"]):
                script.decompose()
            
            # Get text and clean it up
            text = soup.get_text()
            lines = (line.strip() for line in text.splitlines())
            chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
            text = '\n'.join(chunk for chunk in chunks if chunk)
            
            return text[:2000]  # Limit to first 2000 characters
        except Exception as e:
            print(f"❌ Error extracting text from HTML: {str(e)}")
            return ""
    
    def _extract_docx_content(self, file_content: bytes) -> str:
        """
        Extract text from DOCX files with improved error handling
        """
        try:
            # First check if this is actually a valid ZIP file (DOCX format)
            import zipfile
            try:
                with zipfile.ZipFile(io.BytesIO(file_content)) as test_zip:
                    # Check if it has the basic DOCX structure
                    if 'word/document.xml' not in test_zip.namelist():
                        raise zipfile.BadZipFile("Not a valid DOCX file structure")
            except zipfile.BadZipFile:
                raise Exception("File is not a valid DOCX format (not a ZIP archive or missing required files)")
            
            # If we get here, it should be a valid DOCX file
            doc = Document(io.BytesIO(file_content))
            paragraphs = []
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    paragraphs.append(paragraph.text.strip())
            
            extracted_text = '\n\n'.join(paragraphs)
            if not extracted_text.strip():
                return "[Empty DOCX document - no text content found]"
            
            return extracted_text
        except Exception as e:
            raise Exception(f"Failed to extract DOCX content: {str(e)}")
    
    def _extract_pdf_content(self, file_content: bytes) -> str:
        """
        Extract text from PDF files with improved error handling
        """
        try:
            # Check if this looks like PDF content
            if not file_content.startswith(b'%PDF'):
                raise Exception("File does not appear to be a valid PDF (missing PDF header)")
            
            pdf_reader = PdfReader(io.BytesIO(file_content))
            text_parts = []
            for page in pdf_reader.pages:
                page_text = page.extract_text()
                if page_text.strip():
                    text_parts.append(page_text.strip())
            
            extracted_text = '\n\n'.join(text_parts)
            if not extracted_text.strip():
                return "[PDF document - no extractable text content found]"
                
            return extracted_text
        except Exception as e:
            raise Exception(f"Failed to extract PDF content: {str(e)}")
    
    def _extract_pptx_content(self, file_content: bytes) -> str:
        """Extract text from PowerPoint files"""
        try:
            presentation = Presentation(io.BytesIO(file_content))
            text_parts = []
            for slide_num, slide in enumerate(presentation.slides, 1):
                slide_texts = []
                for shape in slide.shapes:
                    if hasattr(shape, 'text') and shape.text.strip():
                        slide_texts.append(shape.text.strip())
                if slide_texts:
                    text_parts.append(f"Slide {slide_num}:\n" + '\n'.join(slide_texts))
            return '\n\n'.join(text_parts)
        except Exception as e:
            raise Exception(f"Failed to extract PPTX content: {str(e)}")
    
    async def insert_media_into_rfp(
        self, 
        rfp_name: str, 
        request: InsertMediaRequest,
        tenant_id: str,
        user_id: str,
        access_token: Optional[str] = None
    ) -> InsertMediaResponse:
        """
        Insert media (image/document) into RFP at specified location
        """
        try:
            # Generate unique insertion ID
            insertion_id = str(uuid.uuid4())
            
            # Determine if we need to extract content (for documents going into sections or appendix)
            extract_content = (request.mediaItem.type == 'document' and 
                             request.insertionOptions.location.type in ['section', 'appendix', 'separate_page'])
            
            # Download and save media to blob storage (and extract content if needed)
            blob_url, extracted_content = await self._download_and_save_to_blob(
                media_url=request.mediaItem.url,
                media_name=request.mediaItem.name,
                tenant_id=tenant_id,
                user_id=user_id,
                extract_content=extract_content,
                access_token=access_token,
                source=request.mediaItem.source,
                media_type=request.mediaItem.type
            )
            
            # Update the media item URL to use blob URL instead of SharePoint URL
            original_url = request.mediaItem.url
            request.mediaItem.url = blob_url
            
            print(f"🔄 Media URL updated from {request.mediaItem.source} to Blob: {original_url} -> {blob_url}")
            
            # Get RFP data from Cosmos DB
            rfp_data = await self.rfp_data_service.get_rfp_by_name(
                tenant_id=tenant_id,
                user_id=user_id,
                rfp_name=rfp_name
            )
            
            if not rfp_data:
                return InsertMediaResponse(
                    success=False,
                    error=ErrorInfo(
                        code="RFP_NOT_FOUND",
                        message=f"RFP '{rfp_name}' not found for the user"
                    )
                )
            
            # Validate media item (supporting images and documents)
            if request.mediaItem.type not in ['image', 'document']:
                return InsertMediaResponse(
                    success=False,
                    error=ErrorInfo(
                        code="MEDIA_TYPE_NOT_SUPPORTED",
                        message=f"Media type '{request.mediaItem.type}' is not supported",
                        details=ErrorDetails(
                            field="mediaItem.type",
                            reason=f"Only 'image' and 'document' types are supported"
                        )
                    )
                )
            
            # Handle different location types
            if request.insertionOptions.location.type == 'section':
                result = await self._insert_into_section(
                    rfp_data, request, insertion_id, tenant_id, user_id, extracted_content
                )
            elif request.insertionOptions.location.type == 'separate_page':
                result = await self._insert_as_separate_page(
                    rfp_data, request, insertion_id, tenant_id, user_id, extracted_content
                )
            elif request.insertionOptions.location.type == 'appendix':
                result = await self._insert_into_appendix(
                    rfp_data, request, insertion_id, tenant_id, user_id, extracted_content
                )
            else:
                return InsertMediaResponse(
                    success=False,
                    error=ErrorInfo(
                        code="INVALID_LOCATION_TYPE",
                        message=f"Location type '{request.insertionOptions.location.type}' is not supported"
                    )
                )
            
            return result
            
        except Exception as e:
            return InsertMediaResponse(
                success=False,
                error=ErrorInfo(
                    code="INTERNAL_ERROR",
                    message=f"Failed to insert media: {str(e)}"
                )
            )
    
    async def _insert_into_section(
        self, 
        rfp_data: Dict, 
        request: InsertMediaRequest, 
        insertion_id: str,
        tenant_id: str,
        user_id: str,
        extracted_content: Optional[str] = None
    ) -> InsertMediaResponse:
        """Insert media into a specific section"""
        sections = rfp_data.get("sections", [])
        target_section = None
        section_index = None
        
        print(f"🔍 Looking for section. Request sectionId: '{request.insertionOptions.location.sectionId}', sectionTitle: '{request.insertionOptions.location.sectionTitle}'")
        print(f"📋 Available sections in RFP:")
        for idx, section in enumerate(sections):
            section_name = section.get("SectionName", "")
            print(f"  {idx}: '{section_name}'")
        
        # Find target section by ID or title
        for idx, section in enumerate(sections):
            section_name = section.get("SectionName", "")
            if (request.insertionOptions.location.sectionId and 
                section_name == request.insertionOptions.location.sectionId) or \
               (request.insertionOptions.location.sectionTitle and 
                section_name.lower() == request.insertionOptions.location.sectionTitle.lower()):
                target_section = section
                section_index = idx
                print(f"✅ Found matching section at index {idx}: '{section_name}'")
                break
        
        if not target_section:
            print(f"❌ Target section not found!")
            return InsertMediaResponse(
                success=False,
                error=ErrorInfo(
                    code="SECTION_NOT_FOUND",
                    message="Target section not found in RFP",
                    details=ErrorDetails(
                        field="insertionOptions.location.sectionId",
                        reason="Section ID or title does not exist in the RFP"
                    )
                )
            )
        
        # Handle different media types
        if request.mediaItem.type == 'image':
            # Handle image insertion
            rendered_html = self._generate_image_html(request)
            print(f"🎨 Generated HTML: {rendered_html[:100]}...")
            
            print(f"🖼️ Setting image URL in Image field: {request.mediaItem.url}")
            
            # Create image object with new format
            image_object = {
                "url": request.mediaItem.url,
                "isInline": request.insertionOptions.displayOptions.type == 'inline',
                "caption": request.insertionOptions.displayOptions.caption or request.mediaItem.name
            }
            
            # Update Image field with array format
            current_images = sections[section_index].get("Image", [])
            if isinstance(current_images, str):  # Handle legacy string format
                current_images = [] if not current_images else [{"url": current_images, "isInline": False, "caption": ""}]
            elif not isinstance(current_images, list):
                current_images = []
            
            current_images.append(image_object)
            sections[section_index]["Image"] = current_images
            
            print(f"💾 Attempting to update section '{target_section.get('SectionName', '')}' in RFP '{rfp_data.get('RFP_Name', '')}'")
            
            # Update RFP in Cosmos DB - pass the new image array format
            update_success = await self.rfp_data_service.update_section_content(
                tenant_id=tenant_id,
                user_id=user_id,
                rfp_name=rfp_data.get("RFP_Name", ""),
                section_name=target_section.get("SectionName", ""),
                new_content=target_section.get("Content", ""),  # Keep original content unchanged
                image_data=sections[section_index]["Image"]  # Pass the new image array
            )
            
        elif request.mediaItem.type == 'document':
            # Handle document content insertion
            print(f"📄 Processing document content insertion")
            print(f"📄 Extracted content available: {extracted_content is not None}")
            print(f"📄 Extracted content length: {len(extracted_content) if extracted_content else 0}")
            
            if not extracted_content:
                print(f"❌ No extracted content available for document: {request.mediaItem.name}")
                return InsertMediaResponse(
                    success=False,
                    error=ErrorInfo(
                        code="CONTENT_EXTRACTION_FAILED",
                        message="Failed to extract content from document"
                    )
                )
            
            # Get current content and append document content
            current_content = target_section.get("Content", "")
            
            # Add document content with a separator and header
            document_header = f"\n\n--- Content from {request.mediaItem.name} ---\n"
            new_content = current_content + document_header + extracted_content
            
            print(f"📝 Appending {len(extracted_content)} characters to section content")
            print(f"📝 New total content length: {len(new_content)} characters")
            
            # Update section content in the sections array
            sections[section_index]["Content"] = new_content
            
            # Update RFP in Cosmos DB with new content
            update_success = await self.rfp_data_service.update_section_content(
                tenant_id=tenant_id,
                user_id=user_id,
                rfp_name=rfp_data.get("RFP_Name", ""),
                section_name=target_section.get("SectionName", ""),
                new_content=new_content,
                image_data=target_section.get("Image", [])  # Keep existing images
            )
            
            # Generate HTML for response (document link)
            rendered_html = f'<div class="document-reference"><a href="{request.mediaItem.url}" target="_blank">{request.mediaItem.name}</a><br><small>Content appended to section</small></div>'
        
        else:
            return InsertMediaResponse(
                success=False,
                error=ErrorInfo(
                    code="UNSUPPORTED_MEDIA_TYPE",
                    message=f"Media type '{request.mediaItem.type}' is not supported"
                )
            )
        
        print(f"💾 Update result: {update_success}")
        
        if not update_success:
            return InsertMediaResponse(
                success=False,
                error=ErrorInfo(
                    code="UPDATE_FAILED",
                    message="Failed to update section with media"
                )
            )
        
        # Create response
        return InsertMediaResponse(
            success=True,
            data=InsertMediaResponseData(
                insertionId=insertion_id,
                mediaItem=MediaItemResponse(
                    id=request.mediaItem.id,
                    name=request.mediaItem.name,
                    type=request.mediaItem.type,
                    url=request.mediaItem.url,
                    insertedAt=datetime.utcnow().isoformat()
                ),
                location=LocationData(
                    type="section",
                    sectionId=target_section.get("SectionName"),
                    sectionTitle=target_section.get("SectionName"),
                    position=len(sections)
                ),
                displayInfo=DisplayInfo(
                    type=request.insertionOptions.displayOptions.type,
                    caption=request.insertionOptions.displayOptions.caption,
                    figureNumber=request.insertionOptions.displayOptions.figureNumber,
                    renderedHtml=rendered_html
                )
            ),
            message="Media successfully inserted into section"
        )
    
    async def _insert_as_separate_page(
        self, 
        rfp_data: Dict, 
        request: InsertMediaRequest, 
        insertion_id: str,
        tenant_id: str,
        user_id: str,
        extracted_content: Optional[str] = None
    ) -> InsertMediaResponse:
        """Insert media as a separate page/section"""
        sections = rfp_data.get("sections", [])
        
        # Create new section for the media
        new_section_name = f"Media - {request.mediaItem.name}"
        
        # Handle different media types
        if request.mediaItem.type == 'image':
            # Handle image insertion as separate page
            rendered_html = self._generate_image_html(request)
            
            # Create image object with new format
            image_object = {
                "url": request.mediaItem.url,
                "isInline": request.insertionOptions.displayOptions.type == 'inline',
                "caption": request.insertionOptions.displayOptions.caption or request.mediaItem.name
            }
            
            new_section = {
                "SectionName": new_section_name,
                "Content": "",  # Don't add HTML content for images
                "Image": [image_object],  # Use new array format
                "Confidence": 100
            }
            
        elif request.mediaItem.type == 'document':
            # Handle document content insertion as separate page
            print(f"📄 Processing document content insertion as separate page")
            print(f"📄 Extracted content available: {extracted_content is not None}")
            print(f"📄 Extracted content length: {len(extracted_content) if extracted_content else 0}")
            
            if not extracted_content:
                print(f"❌ No extracted content available for document: {request.mediaItem.name}")
                return InsertMediaResponse(
                    success=False,
                    error=ErrorInfo(
                        code="CONTENT_EXTRACTION_FAILED",
                        message="Failed to extract content from document"
                    )
                )
            
            # Create section with document content
            document_header = f"--- Content from {request.mediaItem.name} ---\n\n"
            new_content = document_header + extracted_content
            
            new_section = {
                "SectionName": new_section_name,
                "Content": new_content,
                "Image": [],  # No images for document content
                "Confidence": 100
            }
            
            print(f"📝 Creating separate page with {len(extracted_content)} characters of content")
            
            # Generate HTML for response (document link)
            rendered_html = f'<div class="document-reference"><a href="{request.mediaItem.url}" target="_blank">{request.mediaItem.name}</a><br><small>Content displayed as separate page</small></div>'
        
        else:
            return InsertMediaResponse(
                success=False,
                error=ErrorInfo(
                    code="UNSUPPORTED_MEDIA_TYPE",
                    message=f"Media type '{request.mediaItem.type}' is not supported"
                )
            )
        
        # Add to sections list
        sections.append(new_section)
        
        # Update RFP in Cosmos DB
        update_success = await self.rfp_data_service.add_new_section(
            tenant_id=tenant_id,
            user_id=user_id,
            rfp_name=rfp_data.get("RFP_Name", ""),
            section_data=new_section
        )
        
        if not update_success:
            return InsertMediaResponse(
                success=False,
                error=ErrorInfo(
                    code="UPDATE_FAILED",
                    message="Failed to create new section for media"
                )
            )
        
        return InsertMediaResponse(
            success=True,
            data=InsertMediaResponseData(
                insertionId=insertion_id,
                mediaItem=MediaItemResponse(
                    id=request.mediaItem.id,
                    name=request.mediaItem.name,
                    type=request.mediaItem.type,
                    url=request.mediaItem.url,
                    insertedAt=datetime.utcnow().isoformat()
                ),
                location=LocationData(
                    type="separate_page",
                    sectionId=new_section_name,
                    sectionTitle=new_section_name,
                    position=len(sections)
                ),
                displayInfo=DisplayInfo(
                    type=request.insertionOptions.displayOptions.type,
                    caption=request.insertionOptions.displayOptions.caption,
                    figureNumber=request.insertionOptions.displayOptions.figureNumber,
                    renderedHtml=rendered_html
                )
            ),
            message="Media successfully inserted as separate page"
        )
    
    async def _insert_into_appendix(
        self, 
        rfp_data: Dict, 
        request: InsertMediaRequest, 
        insertion_id: str,
        tenant_id: str,
        user_id: str,
        extracted_content: Optional[str] = None
    ) -> InsertMediaResponse:
        """Insert media into appendix section"""
        sections = rfp_data.get("sections", [])
        
        # Find or create appendix section
        appendix_section = None
        appendix_index = None
        
        for idx, section in enumerate(sections):
            section_name = section.get("SectionName", "").lower()
            if "appendix" in section_name or "attachment" in section_name:
                appendix_section = section
                appendix_index = idx
                break
        
        # Handle different media types
        if request.mediaItem.type == 'image':
            # Handle image insertion into appendix
            rendered_html = self._generate_image_html(request)
            
            # Create image object with new format
            image_object = {
                "url": request.mediaItem.url,
                "isInline": request.insertionOptions.displayOptions.type == 'inline',
                "caption": request.insertionOptions.displayOptions.caption or request.mediaItem.name
            }
            
            if appendix_section:
                # Add to existing appendix - don't modify content, just add image
                current_images = appendix_section.get("Image", [])
                if isinstance(current_images, str):  # Handle legacy string format
                    current_images = [] if not current_images else [{"url": current_images, "isInline": False, "caption": ""}]
                elif not isinstance(current_images, list):
                    current_images = []
                
                current_images.append(image_object)
                sections[appendix_index]["Image"] = current_images
                
                section_name = appendix_section.get("SectionName", "")
                new_content = appendix_section.get("Content", "")
            else:
                # Create new appendix section
                section_name = "Appendix"
                new_content = ""
                new_appendix = {
                    "SectionName": section_name,
                    "Content": new_content,
                    "Image": [image_object],  # Use new array format
                    "Confidence": 100
                }
                sections.append(new_appendix)
                appendix_index = len(sections) - 1
                
        elif request.mediaItem.type == 'document':
            # Handle document content insertion into appendix
            print(f"📄 Processing document content insertion into appendix")
            print(f"📄 Extracted content available: {extracted_content is not None}")
            print(f"📄 Extracted content length: {len(extracted_content) if extracted_content else 0}")
            
            if not extracted_content:
                print(f"❌ No extracted content available for document: {request.mediaItem.name}")
                return InsertMediaResponse(
                    success=False,
                    error=ErrorInfo(
                        code="CONTENT_EXTRACTION_FAILED",
                        message="Failed to extract content from document"
                    )
                )
            
            # Add document content with a separator and header
            document_header = f"\n\n--- Content from {request.mediaItem.name} ---\n"
            
            if appendix_section:
                # Get current content and append document content
                current_content = appendix_section.get("Content", "")
                new_content = current_content + document_header + extracted_content
                sections[appendix_index]["Content"] = new_content
                section_name = appendix_section.get("SectionName", "")
            else:
                # Create new appendix section with document content
                section_name = "Appendix"
                new_content = document_header + extracted_content
                new_appendix = {
                    "SectionName": section_name,
                    "Content": new_content,
                    "Image": [],
                    "Confidence": 100
                }
                sections.append(new_appendix)
                appendix_index = len(sections) - 1
            
            print(f"📝 Appending {len(extracted_content)} characters to appendix content")
            print(f"📝 New total content length: {len(new_content)} characters")
            
            # Generate HTML for response (document link)
            rendered_html = f'<div class="document-reference"><a href="{request.mediaItem.url}" target="_blank">{request.mediaItem.name}</a><br><small>Content appended to appendix</small></div>'
        
        else:
            return InsertMediaResponse(
                success=False,
                error=ErrorInfo(
                    code="UNSUPPORTED_MEDIA_TYPE",
                    message=f"Media type '{request.mediaItem.type}' is not supported"
                )
            )
        
        # Update RFP in Cosmos DB
        if appendix_section:
            update_success = await self.rfp_data_service.update_section_content(
                tenant_id=tenant_id,
                user_id=user_id,
                rfp_name=rfp_data.get("RFP_Name", ""),
                section_name=section_name,
                new_content=new_content,  # Use the updated content (for documents) or original content (for images)
                image_data=sections[appendix_index].get("Image", [])  # Pass image array
            )
        else:
            update_success = await self.rfp_data_service.add_new_section(
                tenant_id=tenant_id,
                user_id=user_id,
                rfp_name=rfp_data.get("RFP_Name", ""),
                section_data=sections[appendix_index]
            )
        
        if not update_success:
            return InsertMediaResponse(
                success=False,
                error=ErrorInfo(
                    code="UPDATE_FAILED",
                    message="Failed to update appendix with media"
                )
            )
        
        return InsertMediaResponse(
            success=True,
            data=InsertMediaResponseData(
                insertionId=insertion_id,
                mediaItem=MediaItemResponse(
                    id=request.mediaItem.id,
                    name=request.mediaItem.name,
                    type=request.mediaItem.type,
                    url=request.mediaItem.url,
                    insertedAt=datetime.utcnow().isoformat()
                ),
                location=LocationData(
                    type="appendix",
                    sectionId=section_name,
                    sectionTitle=section_name,
                    position=appendix_index + 1
                ),
                displayInfo=DisplayInfo(
                    type=request.insertionOptions.displayOptions.type,
                    caption=request.insertionOptions.displayOptions.caption,
                    figureNumber=request.insertionOptions.displayOptions.figureNumber,
                    renderedHtml=rendered_html
                )
            ),
            message="Media successfully inserted into appendix"
        )
    
    def _generate_image_html(self, request: InsertMediaRequest) -> str:
        """Generate HTML for image insertion"""
        media_item = request.mediaItem
        display_options = request.insertionOptions.displayOptions
        
        # Escape HTML attributes
        name = html.escape(media_item.name)
        url = html.escape(media_item.url)
        
        if display_options.type == 'inline':
            html_content = f'<img src="{url}" alt="{name}" style="max-width: 100%; height: auto;" />'
            if display_options.caption:
                caption = html.escape(display_options.caption)
                html_content += f'\n<p><em>{caption}</em></p>'
        else:  # figure
            figure_num = display_options.figureNumber or 1
            caption = display_options.caption or f"Figure {figure_num}: {name}"
            caption_escaped = html.escape(caption)
            
            html_content = f'''<div class="figure">
    <img src="{url}" alt="{name}" style="max-width: 100%; height: auto;" />
    <p class="caption"><strong>Figure {figure_num}:</strong> {caption_escaped}</p>
</div>'''
        
        return html_content